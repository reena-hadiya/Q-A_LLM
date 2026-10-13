
import os
import re
import logging

import numpy as np
import pymupdf
from docx import Document
from PIL import Image
from dotenv import load_dotenv

from google import genai
from google.genai import types

load_dotenv()

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")

if not GEMINI_API_KEY:
    raise ValueError(
        "GEMINI_API_KEY not found. "
        "Add it to your .env file or environment variables."
    )

client = genai.Client(api_key=GEMINI_API_KEY)

# Primary model followed by backup models.
# You can add or remove models in this list.
GENERATION_MODELS = [
    "gemini-3.8-flash",
    "gemini-3.6-flash",
    "gemini-3.5-flash-lite",
    "gemini-3.1-flash-lite",
]

# Keep embeddings separate from answer-generation models.
EMBEDDING_MODEL = "gemini-embedding-001"
EMBEDDING_DIMENSIONS = 768

CHUNK_SIZE = 2500
CHUNK_OVERLAP = 300

doc_chunks = []
page_metadata = []
chunk_embeddings = []

current_file = None
current_file_type = None


def is_fallback_error(error):
    """
    Return True only for errors where trying another
    generation model may reasonably help.
    """
    message = str(error).lower()

    # Temporary service, network, timeout and quota errors.
    temporary_patterns = [
        r"\b429\b",
        r"\b500\b",
        r"\b502\b",
        r"\b503\b",
        r"\b504\b",
        "resource_exhausted",
        "unavailable",
        "service unavailable",
        "internal server error",
        "deadline exceeded",
        "timed out",
        "timeout",
        "connection reset",
        "connection error",
        "temporarily unavailable",
        "overloaded",
        "rate limit",
    ]

    if any(re.search(pattern, message) for pattern in temporary_patterns):
        return True

    # A particular model might not be available to the API key.
    # Try the next configured model in this case too.
    model_unavailable_patterns = [
        "model not found",
        "model is not found",
        "not found for api version",
        "is not supported for generatecontent",
        "unsupported model",
        "unknown model",
    ]

    return any(
        pattern in message
        for pattern in model_unavailable_patterns
    )


def generate_with_fallback(contents):
    """
    Try generation models in order.
    Return the first successful response.

    Authentication errors, invalid requests and other
    non-retryable errors are raised instead of being hidden.
    """
    errors = []

    for model_name in GENERATION_MODELS:
        try:
            logger.info("Trying Gemini generation model: %s", model_name)

            response = client.models.generate_content(
                model=model_name,
                contents=contents,
            )

            if not response.text:
                raise RuntimeError(
                    f"{model_name} returned an empty response."
                )

            logger.info("Generation succeeded with model: %s", model_name)
            return response

        except Exception as exc:
            errors.append(f"{model_name}: {exc}")

            if not is_fallback_error(exc):
                logger.exception(
                    "Non-retryable error from model %s",
                    model_name,
                )
                raise

            logger.warning(
                "Model %s failed; checking the next model. Error: %s",
                model_name,
                exc,
            )

    error_details = "\n".join(errors)

    raise RuntimeError(
        "All configured Gemini generation models failed. "
        "Check model availability, API quota and service status.\n"
        f"{error_details}"
    )


def split_text(text):
    """Split long text into overlapping chunks."""

    text = (text or "").strip()

    if not text:
        return []

    chunks = []
    start = 0

    while start < len(text):
        end = min(start + CHUNK_SIZE, len(text))

        if end < len(text):
            split_at = text.rfind(" ", start, end)

            if split_at > start:
                end = split_at

        chunk = text[start:end].strip()

        if chunk:
            chunks.append(chunk)

        if end >= len(text):
            break

        start = max(end - CHUNK_OVERLAP, start + 1)

    return chunks


def get_embedding(text, task_type="RETRIEVAL_DOCUMENT"):
    """Generate a text embedding using the dedicated embedding model."""

    if not text or not text.strip():
        raise ValueError("Cannot create an embedding for empty text.")

    result = client.models.embed_content(
        model=EMBEDDING_MODEL,
        contents=text,
        config=types.EmbedContentConfig(
            task_type=task_type,
            output_dimensionality=EMBEDDING_DIMENSIONS,
        ),
    )

    if not result.embeddings or not result.embeddings[0].values:
        raise ValueError("Gemini returned an empty embedding.")

    return np.asarray(
        result.embeddings[0].values,
        dtype=np.float32,
    )


def cosine_similarity(v1, v2):
    """Calculate cosine similarity safely."""

    v1 = np.asarray(v1, dtype=np.float32)
    v2 = np.asarray(v2, dtype=np.float32)

    denominator = np.linalg.norm(v1) * np.linalg.norm(v2)

    if denominator == 0:
        return 0.0

    return float(np.dot(v1, v2) / denominator)


def extract_pdf(file_path):
    chunks = []
    metadata = []

    with pymupdf.open(file_path) as doc:
        for page_num, page in enumerate(doc):
            text = page.get_text("text").strip()

            if not text:
                text = (
                    f"[Scanned or visual page {page_num + 1}. "
                    "Text could not be extracted.]"
                )

            page_chunks = split_text(text)

            for chunk in page_chunks:
                chunks.append(chunk)

                metadata.append({
                    "page_number": page_num + 1,
                    "page_index": page_num,
                })

    return chunks, metadata


def extract_txt(file_path):
    with open(
        file_path,
        "r",
        encoding="utf-8",
        errors="ignore",
    ) as file:
        text = file.read()

    chunks = split_text(text)

    metadata = [
        {"page_number": 1, "page_index": 0}
        for _ in chunks
    ]

    return chunks, metadata


def extract_docx(file_path):
    document = Document(file_path)

    text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    )

    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(
                cell.text.strip() for cell in row.cells
            )

            if row_text.strip():
                text += "\n" + row_text

    chunks = split_text(text)

    metadata = [
        {"page_number": 1, "page_index": 0}
        for _ in chunks
    ]

    return chunks, metadata


def load_document(file_path):
    global doc_chunks
    global page_metadata
    global chunk_embeddings
    global current_file
    global current_file_type

    if not os.path.isfile(file_path):
        raise FileNotFoundError(
            f"Uploaded document not found: {file_path}"
        )

    extension = os.path.splitext(file_path)[1].lower()

    # Reset previous document state before loading another file.
    doc_chunks = []
    page_metadata = []
    chunk_embeddings = []
    current_file = None
    current_file_type = None

    if extension == ".pdf":
        doc_chunks, page_metadata = extract_pdf(file_path)
        current_file_type = "pdf"

    elif extension == ".txt":
        doc_chunks, page_metadata = extract_txt(file_path)
        current_file_type = "txt"

    elif extension == ".docx":
        doc_chunks, page_metadata = extract_docx(file_path)
        current_file_type = "docx"

    elif extension in [".png", ".jpg", ".jpeg"]:
        # Images are processed directly by the generation model.
        doc_chunks = ["[Image Document]"]

        page_metadata = [{
            "page_number": 1,
            "page_index": 0,
        }]

        current_file_type = "image"

    else:
        raise ValueError(
            "Unsupported file type. "
            "Please upload PDF, TXT, DOCX, PNG, JPG or JPEG."
        )

    if current_file_type == "image":
        current_file = file_path
        return len(doc_chunks)

    if not doc_chunks:
        raise ValueError(
            "No readable text found in the uploaded document."
        )

    # Create all embeddings using the same embedding model.
    embeddings = []

    for index, chunk in enumerate(doc_chunks):
        try:
            embedding = get_embedding(
                chunk,
                task_type="RETRIEVAL_DOCUMENT",
            )

            embeddings.append(embedding)

        except Exception as exc:
            doc_chunks = []
            page_metadata = []
            chunk_embeddings = []
            current_file = None
            current_file_type = None

            raise RuntimeError(
                f"Could not create embedding for text chunk "
                f"{index + 1}. Check Gemini embedding model access, "
                f"API quota and model availability. Details: {exc}"
            ) from exc

    chunk_embeddings = np.vstack(embeddings)
    current_file = file_path

    return len(doc_chunks)


def search_document(query):
    if not doc_chunks or current_file is None:
        raise ValueError("No document has been uploaded.")

    if not query or not query.strip():
        raise ValueError("Please enter a question.")

    if current_file_type == "image":
        return {
            "text": "",
            "metadata": page_metadata[0],
            "score": 1.0,
        }

    if len(chunk_embeddings) != len(doc_chunks):
        raise ValueError(
            "Document embeddings are missing. "
            "Please upload the document again."
        )

    # Embedding search intentionally uses the dedicated embedding model.
    query_embedding = get_embedding(
        query,
        task_type="RETRIEVAL_QUERY",
    )

    scores = [
        cosine_similarity(query_embedding, embedding)
        for embedding in chunk_embeddings
    ]

    best_index = int(np.argmax(scores))

    return {
        "text": doc_chunks[best_index],
        "metadata": page_metadata[best_index],
        "score": float(scores[best_index]),
    }


def ask_document(question):
    if not current_file:
        raise ValueError("Please upload a document first.")

    if not question or not question.strip():
        raise ValueError("Please enter a question.")

    # IMAGE QUESTION-ANSWERING
    if current_file_type == "image":
        try:
            with Image.open(current_file) as source_image:
                image = source_image.convert("RGB")

            prompt = f"""
You are a document question-answering assistant.

Analyze the uploaded image and answer the user's question.

Question:
{question}

Rules:
- Answer using only information visible in the image.
- If the answer is not available, say:
  "I don't know based on the uploaded document."
- Do not invent information.
"""

            # Automatically tries the configured models in order.
            response = generate_with_fallback(
                contents=[prompt, image]
            )

            answer = response.text or (
                "I don't know based on the uploaded document."
            )

            return {
                "answer": answer,
                "page": 1,
                "score": 1.0,
            }

        except Exception as exc:
            raise RuntimeError(
                f"Could not analyze the image: {exc}"
            ) from exc

    # TEXT DOCUMENT QUESTION-ANSWERING
    match = search_document(question)

    page_number = match["metadata"]["page_number"]
    context = match["text"]

    prompt = f"""
You are a document question-answering assistant.

Answer the user's question using ONLY the provided
document context.

Document Context:
-----------------
{context}
-----------------

User Question:
{question}

Rules:
1. Use only the provided document context.
2. Do not use outside knowledge.
3. Do not invent facts.
4. If the answer is not present in the context, say exactly:
"I don't know based on the uploaded document."
5. Give a clear and simple answer.
6. Mention the page number when useful.
"""

    try:
        # Automatically tries the configured models in order.
        response = generate_with_fallback(contents=prompt)

        answer = response.text or (
            "I don't know based on the uploaded document."
        )

        return {
            "answer": answer,
            "page": page_number,
            "score": match["score"],
        }

    except Exception as exc:
        raise RuntimeError(
            f"Could not generate an answer: {exc}"
        ) from exc







# import os
# import numpy as np
# import pymupdf
# from docx import Document
# from PIL import Image
# from dotenv import load_dotenv

# from google import genai
# from google.genai import types

# load_dotenv()

# GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")

# if not GEMINI_API_KEY:
#     raise ValueError(
#         "GEMINI_API_KEY not found. "
#         "Add it to your .env file or environment variables."
#     )

# client = genai.Client(api_key=GEMINI_API_KEY)

# MODEL_NAME = "gemini-3.8-flash"

# EMBEDDING_MODEL = "gemini-embedding-001"
# EMBEDDING_DIMENSIONS = 768

# CHUNK_SIZE = 2500
# CHUNK_OVERLAP = 300

# doc_chunks = []
# page_metadata = []
# chunk_embeddings = []

# current_file = None
# current_file_type = None

# def split_text(text):
#     """Split long text into overlapping chunks."""

#     text = (text or "").strip()

#     if not text:
#         return []

#     chunks = []
#     start = 0

#     while start < len(text):
#         end = min(start + CHUNK_SIZE, len(text))

#         # Prefer splitting at a word boundary
#         if end < len(text):
#             split_at = text.rfind(" ", start, end)

#             if split_at > start:
#                 end = split_at

#         chunk = text[start:end].strip()

#         if chunk:
#             chunks.append(chunk)

#         if end >= len(text):
#             break

#         start = max(end - CHUNK_OVERLAP, start + 1)

#     return chunks

# def get_embedding(text, task_type="RETRIEVAL_DOCUMENT"):
#     """Generate a text embedding using Gemini."""

#     if not text or not text.strip():
#         raise ValueError("Cannot create an embedding for empty text.")

#     result = client.models.embed_content(
#         model=EMBEDDING_MODEL,
#         contents=text,
#         config=types.EmbedContentConfig(
#             task_type=task_type,
#             output_dimensionality=EMBEDDING_DIMENSIONS
#         )
#     )

#     if not result.embeddings or not result.embeddings[0].values:
#         raise ValueError("Gemini returned an empty embedding.")

#     return np.asarray(
#         result.embeddings[0].values,
#         dtype=np.float32
#     )


# def cosine_similarity(v1, v2):
#     """Calculate cosine similarity safely."""

#     v1 = np.asarray(v1, dtype=np.float32)
#     v2 = np.asarray(v2, dtype=np.float32)

#     denominator = np.linalg.norm(v1) * np.linalg.norm(v2)

#     if denominator == 0:
#         return 0.0

#     return float(np.dot(v1, v2) / denominator)


# def extract_pdf(file_path):
#     chunks = []
#     metadata = []

#     with pymupdf.open(file_path) as doc:
#         for page_num, page in enumerate(doc):
#             text = page.get_text("text").strip()

#             if not text:
#                 text = (
#                     f"[Scanned or visual page {page_num + 1}. "
#                     "Text could not be extracted.]"
#                 )

#             page_chunks = split_text(text)

#             for chunk in page_chunks:
#                 chunks.append(chunk)

#                 metadata.append({
#                     "page_number": page_num + 1,
#                     "page_index": page_num
#                 })

#     return chunks, metadata

# def extract_txt(file_path):
#     with open(
#         file_path,
#         "r",
#         encoding="utf-8",
#         errors="ignore"
#     ) as file:
#         text = file.read()

#     chunks = split_text(text)

#     metadata = [
#         {
#             "page_number": 1,
#             "page_index": 0
#         }
#         for _ in chunks
#     ]

#     return chunks, metadata

# def extract_docx(file_path):
#     document = Document(file_path)

#     text = "\n".join(
#         paragraph.text
#         for paragraph in document.paragraphs
#         if paragraph.text.strip()
#     )

#     for table in document.tables:
#         for row in table.rows:
#             row_text = " | ".join(
#                 cell.text.strip() for cell in row.cells
#             )

#             if row_text.strip():
#                 text += "\n" + row_text

#     chunks = split_text(text)

#     metadata = [
#         {
#             "page_number": 1,
#             "page_index": 0
#         }
#         for _ in chunks
#     ]

#     return chunks, metadata

# def load_document(file_path):
#     global doc_chunks
#     global page_metadata
#     global chunk_embeddings
#     global current_file
#     global current_file_type

#     if not os.path.isfile(file_path):
#         raise FileNotFoundError(
#             f"Uploaded document not found: {file_path}"
#         )

#     extension = os.path.splitext(file_path)[1].lower()

#     doc_chunks = []
#     page_metadata = []
#     chunk_embeddings = []
#     current_file = None
#     current_file_type = None

#     if extension == ".pdf":
#         doc_chunks, page_metadata = extract_pdf(file_path)
#         current_file_type = "pdf"

#     elif extension == ".txt":
#         doc_chunks, page_metadata = extract_txt(file_path)
#         current_file_type = "txt"

#     elif extension == ".docx":
#         doc_chunks, page_metadata = extract_docx(file_path)
#         current_file_type = "docx"

#     elif extension in [".png", ".jpg", ".jpeg"]:
#         # Images are processed directly by Gemini.
#         doc_chunks = ["[Image Document]"]

#         page_metadata = [{
#             "page_number": 1,
#             "page_index": 0
#         }]

#         current_file_type = "image"

#     else:
#         raise ValueError(
#             "Unsupported file type. "
#             "Please upload PDF, TXT, DOCX, PNG, JPG or JPEG."
#         )

#     # Keep existing image handling without text embeddings
#     if current_file_type == "image":
#         current_file = file_path
#         return len(doc_chunks)

#     if not doc_chunks:
#         raise ValueError(
#             "No readable text found in the uploaded document."
#         )

#     # Create Gemini embeddings for text chunks
#     embeddings = []

#     for index, chunk in enumerate(doc_chunks):
#         try:
#             embedding = get_embedding(
#                 chunk,
#                 task_type="RETRIEVAL_DOCUMENT"
#             )

#             embeddings.append(embedding)

#         except Exception as exc:
#             # Do not continue with incomplete embeddings
#             doc_chunks = []
#             page_metadata = []
#             chunk_embeddings = []
#             current_file_type = None

#             raise RuntimeError(
#                 f"Could not create embedding for text chunk "
#                 f"{index + 1}. Check Gemini API access, quota "
#                 f"and model availability. Details: {exc}"
#             ) from exc

#     chunk_embeddings = np.vstack(embeddings)
#     current_file = file_path

#     return len(doc_chunks)

# def search_document(query):
#     if not doc_chunks or current_file is None:
#         raise ValueError(
#             "No document has been uploaded."
#         )

#     if not query or not query.strip():
#         raise ValueError("Please enter a question.")

#     if current_file_type == "image":
#         return {
#             "text": "",
#             "metadata": page_metadata[0],
#             "score": 1.0
#         }

#     if len(chunk_embeddings) != len(doc_chunks):
#         raise ValueError(
#             "Document embeddings are missing. "
#             "Please upload the document again."
#         )

#     query_embedding = get_embedding(
#         query,
#         task_type="RETRIEVAL_QUERY"
#     )

#     scores = [
#         cosine_similarity(query_embedding, embedding)
#         for embedding in chunk_embeddings
#     ]

#     best_index = int(np.argmax(scores))

#     return {
#         "text": doc_chunks[best_index],
#         "metadata": page_metadata[best_index],
#         "score": float(scores[best_index])
#     }


# def ask_document(question):
#     if not current_file:
#         raise ValueError(
#             "Please upload a document first."
#         )

#     if not question or not question.strip():
#         raise ValueError("Please enter a question.")

#     # IMAGE QUESTION-ANSWERING
#     if current_file_type == "image":
#         try:
#             with Image.open(current_file) as source_image:
#                 image = source_image.convert("RGB")

#             prompt = f"""
# You are a document question-answering assistant.

# Analyze the uploaded image and answer the user's question.

# Question:
# {question}

# Rules:
# - Answer using only information visible in the image.
# - If the answer is not available, say:
#   "I don't know based on the uploaded document."
# - Do not invent information.
# """

#             response = client.models.generate_content(
#                 model=MODEL_NAME,
#                 contents=[prompt, image]
#             )

#             answer = response.text

#             if not answer:
#                 answer = (
#                     "I don't know based on the uploaded document."
#                 )

#             return {
#                 "answer": answer,
#                 "page": 1,
#                 "score": 1.0
#             }

#         except Exception as exc:
#             raise RuntimeError(
#                 f"Could not analyze the image: {exc}"
#             ) from exc

#     # TEXT DOCUMENT QUESTION-ANSWERING
#     match = search_document(question)

#     page_number = match["metadata"]["page_number"]
#     context = match["text"]

#     prompt = f"""
# You are a document question-answering assistant.

# Answer the user's question using ONLY the provided
# document context.

# Document Context:
# -----------------
# {context}
# -----------------

# User Question:
# {question}

# Rules:
# 1. Use only the provided document context.
# 2. Do not use outside knowledge.
# 3. Do not invent facts.
# 4. If the answer is not present in the context, say exactly:
# "I don't know based on the uploaded document."
# 5. Give a clear and simple answer.
# 6. Mention the page number when useful.
# """

#     try:
#         response = client.models.generate_content(
#             model=MODEL_NAME,
#             contents=prompt
#         )

#         answer = response.text

#         if not answer:
#             answer = (
#                 "I don't know based on the uploaded document."
#             )

#         return {
#             "answer": answer,
#             "page": page_number,
#             "score": match["score"]
#         }

#     except Exception as exc:
#         raise RuntimeError(
#             f"Could not generate an answer: {exc}"
#         ) from exc






# import os
# import numpy as np
# import pymupdf
# from docx import Document
# from PIL import Image
# from dotenv import load_dotenv

# from sentence_transformers import SentenceTransformer
# from google import genai
# from google.genai import types

# load_dotenv()

# GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")

# if not GEMINI_API_KEY:
#     raise ValueError("GEMINI_API_KEY not found in .env file")

# client = genai.Client(api_key=GEMINI_API_KEY)

# MODEL_NAME = "gemini-3.5-flash"

# embedder = SentenceTransformer("all-MiniLM-L6-v2")


# doc_chunks = []
# page_metadata = []
# chunk_embeddings = []

# current_file = None
# current_file_type = None

# def cosine_similarity(v1, v2):

#     return np.dot(v1, v2) / (
#         np.linalg.norm(v1) * np.linalg.norm(v2)
#     )

# def extract_pdf(file_path):

#     chunks = []
#     metadata = []

#     doc = pymupdf.open(file_path)

#     for page_num in range(len(doc)):

#         page = doc[page_num]

#         text = page.get_text("text").strip()

#         if text:

#             chunks.append(text)

#         else:

#             chunks.append(
#                 f"[Scanned/Visual Page {page_num + 1}]"
#             )

#         metadata.append({
#             "page_number": page_num + 1,
#             "page_index": page_num
#         })

#     doc.close()

#     return chunks, metadata

# def extract_txt(file_path):

#     with open(
#         file_path,
#         "r",
#         encoding="utf-8",
#         errors="ignore"
#     ) as file:

#         text = file.read()

#     return [text], [
#         {
#             "page_number": 1,
#             "page_index": 0
#         }
#     ]

# def extract_docx(file_path):

#     document = Document(file_path)

#     text = "\n".join(
#         paragraph.text
#         for paragraph in document.paragraphs
#         if paragraph.text.strip()
#     )

#     return [text], [
#         {
#             "page_number": 1,
#             "page_index": 0
#         }
#     ]

# def load_document(file_path):

#     global doc_chunks
#     global page_metadata
#     global chunk_embeddings
#     global current_file
#     global current_file_type

#     extension = os.path.splitext(file_path)[1].lower()

#     if extension == ".pdf":

#         doc_chunks, page_metadata = extract_pdf(file_path)
#         current_file_type = "pdf"

#     elif extension == ".txt":

#         doc_chunks, page_metadata = extract_txt(file_path)
#         current_file_type = "txt"

#     elif extension == ".docx":

#         doc_chunks, page_metadata = extract_docx(file_path)
#         current_file_type = "docx"

#     elif extension in [".png", ".jpg", ".jpeg"]:

#         doc_chunks = ["[Image Document]"]

#         page_metadata = [
#             {
#                 "page_number": 1,
#                 "page_index": 0
#             }
#         ]

#         current_file_type = "image"

#     else:

#         raise ValueError(
#             "Unsupported file type"
#         )

#     current_file = file_path

#     # Create embeddings for text documents
#     if current_file_type != "image":

#         chunk_embeddings = embedder.encode(
#             doc_chunks,
#             convert_to_numpy=True
#         )

#     else:

#         chunk_embeddings = []

#     return len(doc_chunks)

# def search_document(query):

#     if not doc_chunks:

#         raise ValueError(
#             "No document has been uploaded."
#         )

#     # Images don't use text retrieval
#     if current_file_type == "image":

#         return {
#             "text": "",
#             "metadata": page_metadata[0],
#             "score": 1.0
#         }

#     query_embedding = embedder.encode(
#         query,
#         convert_to_numpy=True
#     )

#     scores = [
#         cosine_similarity(
#             query_embedding,
#             document_embedding
#         )
#         for document_embedding in chunk_embeddings
#     ]

#     best_index = int(np.argmax(scores))

#     return {
#         "text": doc_chunks[best_index],
#         "metadata": page_metadata[best_index],
#         "score": float(scores[best_index])
#     }

# def ask_document(question):

#     if not current_file:

#         raise ValueError(
#             "Please upload a document first."
#         )

#     if current_file_type == "image":

#         image = Image.open(current_file)

#         prompt = f"""
# You are a document question-answering assistant.

# Analyze the uploaded image and answer the user's question.

# Question:
# {question}

# Rules:
# - Answer using only information visible in the image.
# - If the answer is not available in the image, say:
#   "I don't know based on the uploaded document."
# - Do not invent information.
# """

#         response = client.models.generate_content(
#             model=MODEL_NAME,
#             contents=[
#                 prompt,
#                 image
#             ]
#         )

#         return {
#             "answer": response.text,
#             "page": 1,
#             "score": 1.0
#         }

#     match = search_document(question)

#     page_number = match["metadata"]["page_number"]

#     context = match["text"]

#     prompt = f"""
# You are a document question-answering assistant.

# Answer the user's question using ONLY the provided
# document context.

# Document Context:
# -----------------
# {context}
# -----------------

# User Question:
# {question}

# Rules:

# 1. Use only the document context.
# 2. Do not use outside knowledge.
# 3. Do not invent facts.
# 4. If the answer is not present in the context,
#    say exactly:

# "I don't know based on the uploaded document."

# 5. Give a clear and simple answer.
# 6. If useful, mention the page number.
# """

#     response = client.models.generate_content(
#         model=MODEL_NAME,
#         contents=prompt
#     )

#     return {
#         "answer": response.text,
#         "page": page_number,
#         "score": match["score"]
#     }